"""Independent acceptance oracle. Reads an XLSX; never modifies it.

Usage: python3 scripts/verify-workbook.py FILE.xlsx < expected-sheet-rows.json
Requires openpyxl in the explicitly selected validation environment, not the app.
"""
import json
import sys
import warnings
import openpyxl

expected = json.load(sys.stdin)
with warnings.catch_warnings(record=True) as seen:
    warnings.simplefilter("always")
    workbook = openpyxl.load_workbook(sys.argv[1], data_only=False)
    assert not seen, [str(w.message) for w in seen]
assert workbook.sheetnames == list(expected), workbook.sheetnames
for name, rows in expected.items():
    sheet = workbook[name]
    actual = [list(row) for row in sheet.iter_rows(values_only=True)]
    # Empty inline strings may be represented as None by readers.
    normalized = [["" if value is None else value for value in row] for row in actual]
    assert normalized == rows, {"sheet": name, "actual": normalized, "expected": rows}
    assert sheet.freeze_panes == "A2", (name, sheet.freeze_panes)
    assert sheet.auto_filter.ref, name
    for row in sheet:
        for cell in row:
            assert cell.data_type != "f", (name, cell.coordinate, "unexpected formula")
            assert cell.hyperlink is None, (name, cell.coordinate, "unexpected hyperlink")
assert not workbook._external_links
assert workbook.vba_archive is None
workbook.close()
print(json.dumps({"reader": "openpyxl", "version": openpyxl.__version__, "sheets": list(expected), "exactValues": True, "noExecutableCells": True}))
